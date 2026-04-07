from datetime import datetime
from flask import render_template, request, jsonify, Blueprint,current_app, make_response
from flask_login.utils import login_required, current_user
from app import util
from app.model.cms import Category

from app.util.common import get_short_id, time_now_name, upload, get_download_file_path
from app.util.permission import admin_perm
from app.util import R, request_form_auto_fill
from app.util.model import query_model
from app.util.dict_data import get_dict_data, get_label_by_value
from app.ext import db

from openpyxl import Workbook

from app.admin.views.cms import cms_bp
from app.model import Article,Tag
import os

bp = Blueprint('article', __name__)
cms_bp.register_blueprint(bp, url_prefix='/article')
tpl_prefix = 'admin/cms/article/'
model_name = '文章管理'

@bp.route('/',methods=['GET'])
@login_required
@admin_perm(model_name, '查看列表', permission='cms:article:view')
def index():
    """
    列表
    """
    return render_template( tpl_prefix + 'list.html')

@bp.route('/list',methods=['GET','POST'])
@login_required
@admin_perm(model_name, '获取列表数据', permission='cms:article:view')
def list():
    """
    获取列表
    """
    list = query_model(Article, search_query, need_scope=True)
    return jsonify(list)

def search_query(query):
    query = query.filter(Article.deleted == 0)
    value = request.values.get('title','').strip()
    if value != '':
        query = query.filter(Article.title.like('%%%s%%' % value))
    value = request.values.get('publish_time','').strip()
    if value != '':
        query = query.filter(Article.publish_time == value)
    value = request.values.get('author_id','').strip()
    if value != '':
        query = query.filter(Article.author_id == value)
    value = request.values.get('category_id','').strip()
    if value != '':
        # query = query.filter(Article.category_id == value)
        category = Category.query.get(value)
        query = query.filter(Article.category_id.in_(category._children_ids()))
    

    query = query.order_by(Article.sn.desc(),Article.id.desc())
    value = request.values.get('ids','').strip()
    if value != '':
        ids = value.split(',')
        query = query.filter(Article.id.in_(ids))
    return query


@bp.route('/export',methods=['GET','POST'])
@login_required
@admin_perm(model_name, '导出excel', permission='cms:article:export')
def export():
    list = query_model(Article, search_query, is_page=False, need_scope=True)
    
    title = ["备注","机构id","标题","编码","编辑器","md内容","内容","简述","缩略图","状态","访问统计","","发布时间","","","","1=管理员,2=普通账号,3=vip","是否抓取内容","抓取内容的原始url","原作者"]
    state_datas = get_dict_data('cms_article_state')
    wb = Workbook()
    sheet_names = wb.sheetnames
    ws = wb[sheet_names[0]]
    ws.append(title)
    for row in list:
        data = []
        data.append(row.remark)
        data.append(row.org_id)
        data.append(row.title)
        data.append(row.name)
        data.append(row.editor)
        data.append(row.content)
        data.append(row.content_html)
        data.append(row.summary)
        data.append(row.thumbnail)
        data.append(get_label_by_value(state_datas, row.state))
        data.append(row.vc)
        data.append(row.comment_num)
        data.append(row.publish_time)
        data.append(row.author_id)
        data.append(row.category_id)
        data.append(row.h_content)
        data.append(row.h_role)
        data.append(row.is_crawl)
        data.append(row.origin_url)
        data.append(row.origin_author)
        ws.append(data)
    
    abs_filepath, filepath = get_download_file_path("{}.xlsx".format(time_now_name()))
    wb.save(abs_filepath)
    return R.success(msg=filepath)
    

@bp.route('/add',methods=['GET','POST'])
@login_required
@admin_perm(model_name, '添加', permission='cms:article:add')
def add():
    """
    添加
    """
    if request.method == 'POST': # 添加
        o = Article()
        request_form_auto_fill(o)
        if request.files.get('thumbnail'):
            img_url = upload('thumbnail', 'material')
            o.thumbnail = img_url
        if int(o.state) == 1:
            o.publish_time = datetime.now()
        if o.name == '':
            o.name = get_short_id()
        if type(o.publish_time) == str:
            o.publish_time = datetime.strptime(o.publish_time, '%Y-%m-%d %H:%M:%S')
        db.session.add(o)
        db.session.commit()
        if o.name == '':
            o.name = o.id
            db.session.commit()
        return R.success(msg='成功')
    category = None
    category_id = request.args.get('category_id', type=int)
    if category_id:
        category = Category.query.get(category_id)
    return render_template( tpl_prefix + 'add.html', category = category)

@bp.route('/edit', methods=['GET', 'POST'])
@login_required
@admin_perm(model_name, '修改', permission='cms:article:edit')
def edit():
    editor = request.args.get('editor', 'markdown')  # 注意：GET 时也需要 editor

    if request.method == 'POST':
        # === 1. 安全解析 id ===
        id_str = request.form.get('id', '').strip()
        try:
            article_id = int(id_str) if id_str else None
        except (ValueError, TypeError):
            article_id = None

        # === 2. 获取或新建对象 ===
        if article_id:
            o = Article.query.get(article_id)
            if not o:
                return R.error("文章不存在")
            is_new = False
        else:
            o = Article()
            is_new = True

        # === 3. 填充表单数据（包括可能的 id=''）===
        request_form_auto_fill(o)

        # === 4. 【关键】修复 id 字段 ===
        if is_new:
            o.id = None  # 让数据库自增
        else:
            o.id = article_id  # 确保是整数

        # === 5. 校验并处理 category_id ===
        category_id = request.form.get('category_id')
        if category_id:
            try:
                category_id = int(category_id)
                category = Category.query.get(category_id)
                if not category:
                    return R.error("所选分类不存在")
                o.category_id = category_id
            except (ValueError, TypeError):
                return R.error("分类ID格式错误")
        else:
            o.category_id = None

        # === 6. 处理 publish_time（字符串 → datetime）===
        if isinstance(o.publish_time, str) and o.publish_time.strip():
            try:
                o.publish_time = datetime.strptime(o.publish_time.strip(), '%Y-%m-%d %H:%M:%S')
            except ValueError:
                return R.error("发布时间格式错误，应为 YYYY-MM-DD HH:MM:SS")
        elif not o.publish_time:
            o.publish_time = None

        # === 7. Markdown 转 HTML ===
        if o.editor == 'markdown':
            o.content_html, o.toc = o.md_convert()

        # === 8. 自动生成 name（如果为空）===
        if not o.name or o.name.strip() == '':
            if is_new:
                o.name = get_short_id()
            else:
                o.name = str(o.id)

        # === 9. 检查 name 唯一性（排除自己）===
        existing = Article.query.filter(
            Article.name == o.name,
            Article.id != (o.id or 0),
            Article.deleted == 0
        ).first()
        if existing:
            return R.error(f"文章编码 '{o.name}' 已存在")

        # === 10. 保存到数据库 ===
        try:
            if is_new:
                db.session.add(o)
            db.session.flush()  # 获取 o.id（用于 name）

            # 如果 name 还是空（极端情况），用 id 补
            if not o.name:
                o.name = str(o.id)

            # === 11. 处理 tags（正确方式）===
            new_tags = []
            tags_input = request.form.get('tagsinput', '')
            for tag_name in tags_input.split(','):
                tag_name = tag_name.strip()
                if tag_name:
                    t = Tag.add(tag_name)
                    if t:  # 确保返回有效对象
                        new_tags.append(t)
            o.tags = new_tags

            db.session.commit()
            return R.success("保存成功")

        except Exception as e:
            db.session.rollback()
            current_app.logger.error(f"保存文章失败: {e}")
            return R.error("保存失败，请联系管理员")

    # ====== GET 请求：渲染编辑页面 ======
    id = request.args.get('id', type=int)
    o = Article.query.get(id) if id else None
    if o is None:
        o = Article()
        o.vc = 0
        category_id = request.args.get('category_id', type=int)
        if category_id:
            category = Category.query.get(category_id)
            if category:
                setattr(o, 'category_tmp', category)
                o.category_id = category.id

    if not editor:
        editor = o.editor or current_app.config.get('H3BLOG_EDITOR', 'markdown')

    template = 'edit_tinymce.html' if editor == 'tinymce' else 'edit_md.html'
    return render_template(tpl_prefix + template, d=o)

@bp.route('/remove',methods=['GET','POST'])
@login_required
@admin_perm(model_name, '删除', permission='cms:article:remove')
def remove():
    """
    根据id删除
    """
    ids = request.values.get('ids', type=str)
    ids = [int(i) for i in ids.split(',')]
    Article.query.filter(Article.id.in_(ids)).update({'deleted': 1})
    return R.success()


@bp.post('/change_state')
@login_required
@admin_perm(model_name, '修改文章状态')
def change_state():
    """
    修改文章状态
    """
    id = request.form.get('id', 0, type=int)
    state = request.form.get('state', '', type=int)
    a = Article.query.get(id)
    if a:
        a.state = state
        if a.state ==1 and a.publish_time is None:
            a.publish_time = datetime.now()
        db.session.commit()
        return R.success(msg='操作成功')
    return R.error(msg='操作失败')


@bp.route('/import_article',methods=['POST'])
def import_article():
    from readability import Document
    url = request.form.get('url')
    download_img = request.form.get('download_img', 0, type=int)
    is_download_img = True if download_img == 1 else False
    req_html = request.form.get('html')
    if req_html and len(req_html) > 0:
        html = req_html
    else:
        html = util.open_url(url)
        html = util.strdecode(html)
        
        doc = Document(html)
        html = doc.summary()
    markdown = util.html2markdown(html, url, is_download_img, '')
    return jsonify(markdown)

@bp.post('/exchange_sn')
@login_required
@admin_perm(model_name, '交换文章排序')
def exchange_sn():
    """
    交换文章排序
    """
    from_id = request.form.get('from_id', type=int)
    to_id = request.form.get('to_id', type=int)
    from_a = Article.query.get(from_id)
    to_a = Article.query.get(to_id)
    from_sn = from_a.sn
    from_a.sn = to_a.sn
    to_a.sn = from_sn
    return R.success(msg='操作成功')

@bp.post('/change_sn')
@login_required
@admin_perm(model_name, '修改文章排序')
def change_sn():
    """
    修改文章排序
    """
    id = request.form.get('id', type=int)
    sn = request.form.get('sn', type=int)
    article = Article.query.get(id)
    article.sn = sn
    return R.success(msg='操作成功')

@bp.post('/change_category')
@login_required
@admin_perm(model_name, '修改所属栏目')
def change_category():
    """
    修改所属栏目
    """
    ids = request.form.get('ids','')
    ids = [int(id.strip()) for id in ids.split(',') if id.strip() != '']
    category_id = request.form.get('category_id', type=int)
    Article.query.filter(Article.id.in_(ids)).update({'category_id': category_id})
    return R.success(msg='操作成功')